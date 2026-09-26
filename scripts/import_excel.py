import openpyxl,json,re,sys,uuid
from pathlib import Path
p=Path(sys.argv[1] if len(sys.argv)>1 else 'Копия  (1).xlsx'); wb=openpyxl.load_workbook(p,data_only=True)
prices={str(r[0]).strip():float(r[1] or 0) for r in wb['Цены продуктов'].iter_rows(min_row=3,values_only=True) if r[0]}
dishes=[]
for r in wb['Справочники'].iter_rows(min_row=2,values_only=True):
 name=r[0]
 if not name: continue
 sale=float(r[2] or 0) if isinstance(r[2],(int,float)) else 0
 recipe={}
 if r[3]:
  for part in str(r[3]).split('|'):
   if '=' in part:
    n,q=part.rsplit('=',1); recipe[n]=float(q)
 cost=sum(prices.get(n,0)*q for n,q in recipe.items())
 dishes.append({'name':name,'sale_price':sale,'cost_price':cost,'recipe':recipe})
out=Path('supabase/seed'); out.mkdir(parents=True,exist_ok=True); (out/'catalog.json').write_text(json.dumps({'ingredients':[{'name':n,'price':v,'unit':'шт'} for n,v in prices.items()],'dishes':dishes},ensure_ascii=False,indent=2),encoding='utf8'); print(f'ingredients={len(prices)} dishes={len(dishes)}')
